import io
import openpyxl
import pandas as pd

def process_ntuc_batch_po_file(file_bytes_or_buffer, filename=""):
    """
    Directly extracts all sheets/tabs from an NTUC Batch PO workbook.
    Every tab (1, 2, 3...) contains a unique PO No and Store.
    """
    if isinstance(file_bytes_or_buffer, bytes):
        buffer = io.BytesIO(file_bytes_or_buffer)
    else:
        buffer = file_bytes_or_buffer

    # Load workbook with data values evaluated
    wb = openpyxl.load_workbook(buffer, data_only=True)
    all_data = []

    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        
        # Read header cells specific to this sheet
        buyer = ws['B1'].value or "NTUC FAIRPRICE CO-OPERATIVE LIMITED"
        po_no = str(ws['B2'].value or "").strip()
        order_date = ws['B3'].value or ""
        delivery_date = ws['B4'].value or ""
        supplier_name = ws['B6'].value or "ZENXIN AGRI-ORGANIC FOOD PTE LTD"
        supplier_code = str(ws['B7'].value or "39309").strip()

        # Iterate over item rows starting from row 12 (1-based index)
        for r in range(12, ws.max_row + 1):
            s_no = ws.cell(row=r, column=1).value
            if s_no is None:
                continue

            try:
                s_no_clean = int(float(str(s_no).replace(",", "")))
            except (ValueError, TypeError):
                continue

            ean = str(ws.cell(row=r, column=2).value or "").strip()
            desc = str(ws.cell(row=r, column=3).value or "").strip()
            stock_code = str(ws.cell(row=r, column=4).value or "").strip()
            pack_size = ws.cell(row=r, column=5).value or 1
            uom = ws.cell(row=r, column=6).value or "EA"

            # Parse numeric fields
            def to_num(val):
                try:
                    return float(str(val).replace(",", ""))
                except (ValueError, TypeError):
                    return 0.0

            unit_price = to_num(ws.cell(row=r, column=7).value)
            st_wh_no = str(ws.cell(row=r, column=8).value or "").strip()
            st_wh_name = str(ws.cell(row=r, column=9).value or "").strip()
            qty = to_num(ws.cell(row=r, column=10).value)
            amount = to_num(ws.cell(row=r, column=11).value)

            all_data.append({
                "Source File": filename,
                "Sheet / Tab": str(sheetname),
                "PO No": po_no,
                "ST/WH No": st_wh_no,
                "ST/WH Name": st_wh_name,
                "Order Date": order_date,
                "Delivery Date": delivery_date,
                "Buyer Name": buyer,
                "Supplier Name": supplier_name,
                "Supplier Code": supplier_code,
                "S/No": s_no_clean,
                "EAN": ean,
                "Description": desc,
                "Buyer Stock Code": stock_code,
                "Pack Size": pack_size,
                "UOM": uom,
                "Unit Price (SGD)": unit_price,
                "Delivery Quantity": qty,
                "Amount (SGD)": amount
            })

    wb.close()
    return all_data


def combine_all_ntuc_batch_files(uploaded_files):
    """
    Combines all sheets from all uploaded batch PO workbooks.
    """
    all_rows = []
    for file_obj in uploaded_files:
        filename = getattr(file_obj, "name", "")
        file_bytes = file_obj.getvalue() if hasattr(file_obj, "getvalue") else file_obj
        rows = process_ntuc_batch_po_file(file_bytes, filename=filename)
        all_rows.extend(rows)

    if not all_rows:
        return pd.DataFrame()

    return pd.DataFrame(all_rows)
